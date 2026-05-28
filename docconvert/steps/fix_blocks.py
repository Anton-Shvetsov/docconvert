"""Block-level layout on TMP/thesis.docx:
- center paragraphs containing inline images, strip first-line indent;
- zero first-line indent on figure/table captions and on display-equation paragraphs;
- center figure captions; left-align table captions;
- restore the leading space lost from "Рисунок N — title" when crossref versions
  mismatch (best effort);
- strip leftover '(eq:N)' prefixes from cross-references.
"""
from __future__ import annotations

import re

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from ..docx_utils import p_style_id, para_text
from ..pipeline import Context

_MATH_URI = "http://schemas.openxmlformats.org/officeDocument/2006/math"
_OMATH_PARA_TAG = f"{{{_MATH_URI}}}oMathPara"
_EQ_PREFIX_RE = re.compile(r"\(eq:(\d+)\)")
_EMDASH_FIX_RE = re.compile(r"(\d)—")


def _zero_first_line(p_elem):
    pPr = p_elem.find(qn("w:pPr"))
    if pPr is None:
        pPr = OxmlElement("w:pPr")
        p_elem.insert(0, pPr)
    ind = pPr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        pPr.append(ind)
    ind.attrib.pop(qn("w:firstLineChars"), None)
    ind.set(qn("w:firstLine"), "0")
    return pPr


def _set_alignment(p_elem, value: str) -> None:
    pPr = p_elem.find(qn("w:pPr"))
    if pPr is None:
        pPr = OxmlElement("w:pPr")
        p_elem.insert(0, pPr)
    jc = pPr.find(qn("w:jc"))
    if jc is None:
        jc = OxmlElement("w:jc")
        pPr.append(jc)
    jc.set(qn("w:val"), value)


def _zero_spacing(p_elem) -> None:
    pPr = p_elem.find(qn("w:pPr"))
    if pPr is None:
        pPr = OxmlElement("w:pPr")
        p_elem.insert(0, pPr)
    spacing = pPr.find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        pPr.append(spacing)
    spacing.set(qn("w:before"), "0")
    spacing.set(qn("w:after"), "0")


def run(ctx: Context) -> None:
    doc = Document(ctx.paths.working_docx)
    strings = ctx.config.strings

    figure_cap_re = re.compile(rf"^{re.escape(strings.figure_label)}\s+\d")
    table_cap_re = re.compile(rf"^{re.escape(strings.table_label)}\s+\d")

    # Pass A: center figure paragraphs (with inline drawings).
    for para in doc.paragraphs:
        if para._p.find(".//" + qn("w:drawing")) is None:
            continue
        _set_alignment(para._p, "center")
        _zero_first_line(para._p)
        _zero_spacing(para._p)

    # Pass B: caption / equation paragraphs.
    for para in doc.paragraphs:
        sname = (para.style.name or "").lower() if para.style else ""
        sid = p_style_id(para._p) or ""
        sid = sid.lower()
        p_elem = para._p
        text = para_text(p_elem).strip()
        has_drawing = p_elem.find(".//" + qn("w:drawing")) is not None
        has_omath_para = p_elem.find(".//" + _OMATH_PARA_TAG) is not None

        is_caption_style = "caption" in sid or "caption" in sname
        is_table_caption = is_caption_style and table_cap_re.match(text) is not None
        is_figure_caption = (
            not is_table_caption
            and not has_drawing
            and (is_caption_style or bool(figure_cap_re.match(text)))
        )
        is_equation_para = has_omath_para and not has_drawing and not text

        if is_figure_caption or is_table_caption or is_equation_para:
            _zero_first_line(p_elem)

        if is_figure_caption:
            _set_alignment(p_elem, "center")
            for t_el in p_elem.iter(qn("w:t")):
                if t_el.text:
                    t_el.text = _EMDASH_FIX_RE.sub(r"\1 —", t_el.text)

        if is_table_caption:
            _set_alignment(p_elem, "left")

        # Strip leftover (eq:N) → (N)
        for t_el in p_elem.iter(qn("w:t")):
            if t_el.text and "(eq:" in t_el.text:
                t_el.text = _EQ_PREFIX_RE.sub(r"(\1)", t_el.text)

    doc.save(ctx.paths.working_docx)
    print("  blocks: figures centered, caption/equation indents fixed")
