"""Append every appendix listed in ``config.appendices`` to TMP/thesis.docx.

Each entry has:
  - ``code`` — short heading (e.g. "ПРИЛОЖЕНИЕ А"), used to find the heading
    paragraph inside the appendix .docx so its style can be normalised.
  - ``title`` — full descriptive title (used only by insert_toc).
  - ``source`` — file inside input/ (.tex if convert=true, .docx otherwise).
  - ``convert`` — true → run latexpand + pandoc on the .tex source to produce
    a temporary .docx in TMP/; false → use the .docx source as-is.
"""
from __future__ import annotations

import copy
import os
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from ..config import AppendixConfig
from ..pipeline import Context
from ..tools import require, run as run_cmd


def _normalize_heading_style(p_elem, style_id: str = "Heading1") -> None:
    pPr = p_elem.find(qn("w:pPr"))
    if pPr is None:
        pPr = OxmlElement("w:pPr")
        p_elem.insert(0, pPr)
    ps = pPr.find(qn("w:pStyle"))
    if ps is None:
        ps = OxmlElement("w:pStyle")
        pPr.insert(0, ps)
    ps.set(qn("w:val"), style_id)


def _make_pagebreak_paragraph(template_p):
    br_para = copy.deepcopy(template_p)
    for child in list(br_para):
        br_para.remove(child)
    r = OxmlElement("w:r")
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    r.append(br)
    br_para.append(r)
    return br_para


def _convert_tex_to_docx(ctx: Context, source_tex: Path) -> Path:
    """latexpand+pandoc on input/source.tex → TMP/{stem}.docx. Returns the .docx path."""
    stem = source_tex.stem
    flat = ctx.paths.appendix_flat_tex(stem)
    docx_out = ctx.paths.appendix_docx(stem)

    latexpand = require("latexpand")
    pandoc = require("pandoc")

    run_cmd([latexpand, source_tex.name], cwd=ctx.paths.input_dir, stdout_file=flat)

    flat_rel = os.path.relpath(flat, ctx.paths.input_dir)
    out_rel = os.path.relpath(docx_out, ctx.paths.input_dir)
    reference_rel = os.path.relpath(ctx.paths.reference_docx, ctx.paths.input_dir)
    run_cmd(
        [pandoc, flat_rel, "-o", out_rel, f"--reference-doc={reference_rel}"],
        cwd=ctx.paths.input_dir,
    )
    return docx_out


def _force_source_code_font(main_doc) -> None:
    """Force Courier New 10pt on every Source Code paragraph (style definition
    may differ between an appended appendix and the main doc)."""
    for para in main_doc.paragraphs:
        if para.style.name != "Source Code":
            continue
        for run in para.runs:
            run.font.name = "Courier New"
            run.font.size = Pt(10)
        for r_elem in para._element.findall(".//" + qn("w:r")):
            rPr = r_elem.find(qn("w:rPr"))
            if rPr is None:
                rPr = OxmlElement("w:rPr")
                r_elem.insert(0, rPr)
            rFonts = rPr.find(qn("w:rFonts"))
            if rFonts is None:
                rFonts = OxmlElement("w:rFonts")
                rPr.insert(0, rFonts)
            rFonts.set(qn("w:ascii"), "Courier New")
            rFonts.set(qn("w:hAnsi"), "Courier New")
            rFonts.set(qn("w:cs"), "Courier New")
            for short in ("w:sz", "w:szCs"):
                el = rPr.find(qn(short))
                if el is None:
                    el = OxmlElement(short)
                    rPr.append(el)
                el.set(qn("w:val"), "20")


def _append_one(main_doc, appendix_docx: Path, code: str) -> None:
    appendix = Document(appendix_docx)

    # Normalise the heading paragraph so it gets picked up by the TOC.
    heading_para = None
    subtitle_para = None
    for para in appendix.paragraphs:
        if heading_para is None:
            if code.upper() in para.text.upper():
                heading_para = para
                _normalize_heading_style(heading_para._p)
        else:
            if para.text.strip():
                subtitle_para = para
                break
    if heading_para is None:
        print(f"  WARNING: heading '{code}' not found in {appendix_docx.name}")

    if subtitle_para is not None:
        subtitle_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in subtitle_para.runs:
            run.bold = True

    # Page break before the appendix
    last_p = main_doc.paragraphs[-1]._element
    br_para = _make_pagebreak_paragraph(last_p)
    main_doc.element.body.insert(-1, br_para)

    # Copy body elements (skip the source sectPr)
    for elem in appendix.element.body:
        if elem.tag == qn("w:sectPr"):
            continue
        main_doc.element.body.insert(-1, copy.deepcopy(elem))


def run(ctx: Context) -> None:
    appendices = ctx.config.appendices
    if not appendices:
        print("  no appendices configured, skipping")
        return

    main_doc = Document(ctx.paths.working_docx)

    for app in appendices:
        source_path = ctx.paths.appendix_source(app.source)
        if not source_path.exists():
            raise FileNotFoundError(
                f"Appendix source not found: {source_path}. "
                f"Place the {'tex' if app.convert else 'docx'} file in {ctx.paths.input_dir.name}/."
            )

        if app.convert:
            docx_path = _convert_tex_to_docx(ctx, source_path)
        else:
            docx_path = source_path

        print(f"  + {app.code}: {docx_path.name}")
        _append_one(main_doc, docx_path, app.code)

    _force_source_code_font(main_doc)

    main_doc.save(ctx.paths.working_docx)
    print(f"  appendices merged into {ctx.paths.working_docx.name}")
