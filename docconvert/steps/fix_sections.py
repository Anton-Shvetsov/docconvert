"""Section-property fixes on TMP/thesis.docx:
- copy <w:footerReference> elements from the body-level sectPr into every
  intermediate sectPr (pandoc inserts one per \\clearpage but does not copy
  footer references);
- ensure <w:titlePg> is present on every sectPr (so the title page can suppress
  its number);
- set <w:pgNumType w:start="..."> on the first sectPr so numbering begins
  where ``config.page_numbering.start_at`` says.
"""
from __future__ import annotations

import copy

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from ..pipeline import Context


def _ensure_title_pg(sectPr) -> None:
    if sectPr.find(qn("w:titlePg")) is None:
        sectPr.append(OxmlElement("w:titlePg"))


def _set_footer_refs(sectPr, refs) -> None:
    for old in sectPr.findall(qn("w:footerReference")):
        sectPr.remove(old)
    for ref in refs:
        sectPr.append(copy.deepcopy(ref))


def _set_pg_num_start(sectPr, start: int) -> None:
    existing = sectPr.find(qn("w:pgNumType"))
    if existing is not None:
        existing.set(qn("w:start"), str(start))
    else:
        pgn = OxmlElement("w:pgNumType")
        pgn.set(qn("w:start"), str(start))
        sectPr.append(pgn)


def run(ctx: Context) -> None:
    doc = Document(ctx.paths.working_docx)
    body = doc.element.body
    body_sectPr = body.find(qn("w:sectPr"))

    body_refs = []
    if body_sectPr is not None:
        body_refs = [copy.deepcopy(r) for r in body_sectPr.findall(qn("w:footerReference"))]

    for para in doc.paragraphs:
        pPr = para._p.find(qn("w:pPr"))
        if pPr is None:
            continue
        sectPr = pPr.find(qn("w:sectPr"))
        if sectPr is None:
            continue
        _set_footer_refs(sectPr, body_refs)
        _ensure_title_pg(sectPr)

    if body_sectPr is not None:
        _ensure_title_pg(body_sectPr)

    # Page numbering start.
    start_at = ctx.config.page_numbering.start_at
    set_on_first = False
    for para in doc.paragraphs:
        pPr = para._p.find(qn("w:pPr"))
        if pPr is None:
            continue
        sectPr = pPr.find(qn("w:sectPr"))
        if sectPr is None:
            continue
        _set_pg_num_start(sectPr, start_at)
        set_on_first = True
        break
    if not set_on_first and body_sectPr is not None:
        _set_pg_num_start(body_sectPr, start_at)

    doc.save(ctx.paths.working_docx)
    print(f"  sections: footers propagated, page numbering starts at {start_at}")
