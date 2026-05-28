"""Table-specific GOST formatting on TMP/thesis.docx:
- minimum row height (≥ 0.8 cm), single-spaced cells, centered, no first-line indent;
- 14pt space_before on the first non-heading paragraph after each table.
"""
from __future__ import annotations

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from ..pipeline import Context

_MIN_ROW_HEIGHT_TWP = 454  # 0.8 cm in twips (GOST minimum)
_W_TBL = qn("w:tbl")
_W_P = qn("w:p")


def _set_min_row_height(row) -> None:
    trPr = row._tr.find(qn("w:trPr"))
    if trPr is None:
        trPr = OxmlElement("w:trPr")
        row._tr.insert(0, trPr)
    trH = trPr.find(qn("w:trHeight"))
    if trH is None:
        trH = OxmlElement("w:trHeight")
        trPr.append(trH)
    trH.set(qn("w:val"), str(_MIN_ROW_HEIGHT_TWP))
    trH.set(qn("w:hRule"), "atLeast")


def _format_cell_paragraph(p_elem) -> None:
    for r in p_elem.iter(qn("w:r")):
        rPr = r.find(qn("w:rPr"))
        if rPr is None:
            continue
        color = rPr.find(qn("w:color"))
        if color is not None:
            rPr.remove(color)

    pPr = p_elem.find(qn("w:pPr"))
    if pPr is None:
        pPr = OxmlElement("w:pPr")
        p_elem.insert(0, pPr)

    ind = pPr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        pPr.append(ind)
    ind.set(qn("w:firstLine"), "0")

    jc = pPr.find(qn("w:jc"))
    if jc is None:
        jc = OxmlElement("w:jc")
        pPr.append(jc)
    jc.set(qn("w:val"), "center")

    spacing = pPr.find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        pPr.append(spacing)
    spacing.set(qn("w:line"), "240")
    spacing.set(qn("w:lineRule"), "auto")


def _apply_space_before_after_table(p_elem) -> None:
    pPr = p_elem.find(qn("w:pPr"))
    if pPr is None:
        pPr = OxmlElement("w:pPr")
        p_elem.insert(0, pPr)
    spacing = pPr.find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        pPr.append(spacing)
    spacing.attrib.pop(qn("w:beforeLines"), None)
    spacing.attrib.pop(qn("w:beforeAutospacing"), None)
    spacing.set(qn("w:before"), "280")
    ctx_el = pPr.find(qn("w:contextualSpacing"))
    if ctx_el is None:
        ctx_el = OxmlElement("w:contextualSpacing")
        pPr.append(ctx_el)
    ctx_el.set(qn("w:val"), "0")


def run(ctx: Context) -> None:
    doc = Document(ctx.paths.working_docx)

    for table in doc.tables:
        for row in table.rows:
            _set_min_row_height(row)
            for cell in row.cells:
                for para in cell.paragraphs:
                    _format_cell_paragraph(para._p)

    # space_before on the first content paragraph after each table.
    body_children = list(doc.element.body)
    for idx, child in enumerate(body_children):
        if child.tag != _W_TBL:
            continue
        for nxt in body_children[idx + 1:]:
            if nxt.tag == _W_TBL:
                break
            if nxt.tag != _W_P:
                continue
            next_pPr = nxt.find(qn("w:pPr"))
            if next_pPr is not None:
                if next_pPr.find(qn("w:sectPr")) is not None:
                    break
                pStyle_el = next_pPr.find(qn("w:pStyle"))
                if pStyle_el is not None and pStyle_el.get(qn("w:val"), "").lower().startswith("heading"):
                    break
            _apply_space_before_after_table(nxt)
            break

    doc.save(ctx.paths.working_docx)
    print("  tables: row heights, cell formatting, post-table spacing applied")
